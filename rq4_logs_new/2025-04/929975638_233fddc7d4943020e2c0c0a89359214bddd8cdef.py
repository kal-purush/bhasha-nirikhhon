# bulkShortMakerRunner.py

import subprocess
import json
import os
import time
from openpyxl import Workbook, load_workbook
from datetime import datetime
from mutagen.mp3 import MP3


# Constants
EXCEL_FILE = "video_records.xlsx"
SHEET_NAME = "Videos"
SHORTS_JSON_FILE = "bulkShorts_input.json"
OUTPUT_FOLDER = "processed_videos"
# Define columns
columns = [
    'url',
    'video_path',
    'captions_text',
    'youtube_playlist_name',
    'youtube_title',
    'youtube_description',
    'youtube_tags',
    'youtube_channel_name',
    'video_is_about',
    'made_for_kids',
    'schedule_date',
    'youtube_upload_status',
    'last_update_date',
    'created_video_url'
]
# Ensure output folder exists
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# Prepare Excel if it doesn't exist
def setup_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(columns)
        wb.save(EXCEL_FILE)

# Save details in Excel
def save_details_in_excel(video_path, title, description, tags, playlist):
    wb = load_workbook(EXCEL_FILE)
    ws = wb[SHEET_NAME]
    # now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    # row = [video_path, title, description, tags, playlist, 'Pending', now]
    # Prepare row data
    row = [
        "YTShorts",
        video_path,
        description,
        '',
        '',
        '',
        '',
        '',
        '',
        '',
        '',
        'Pending',
        datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        ''
    ]

    ws.append(row)
    wb.save(EXCEL_FILE)
    print(f"✅ Saved record for {video_path}")

def get_total_audio_duration(short_index):
    total_duration = 0.0
    audio_folder = "generated_audio"  # or wherever your audio files are
    
    for i in range(10):  # up to 10 parts
        audio_file = os.path.join(audio_folder, f"short_{short_index}_{i}.mp3")
        if os.path.exists(audio_file):
            audio = MP3(audio_file)
            total_duration += audio.info.length
        else:
            break  # Stop if no more parts
    
    return total_duration

# Main execution
def main():
    # Load shorts data
    with open(SHORTS_JSON_FILE, 'r', encoding='utf-8') as f:
        shorts_data = json.load(f)

    setup_excel()

    for idx, short in enumerate(shorts_data):
        print(f"\n▶️ Recording short {idx + 1}/{len(shorts_data)}")

        output_video_name = f"{OUTPUT_FOLDER}/short_{idx}.mp4"
        video_path = f"short_{idx}"
        # output_video_name = f"short_{idx}.mp4"

        audio_duration = get_total_audio_duration(idx)
        buffer_seconds = 4  # or 5, depending on your subtext time

        recording_duration = audio_duration + buffer_seconds

        # Build Puppeteer command
        cmd = [
            "node", "puppeteer-recorder.js",
            str(idx), str(output_video_name), str(recording_duration)
        ]

        print("Running Puppeteer with:", cmd)
        subprocess.run(cmd)

        # Save details in Excel
        title = short.get('title', '')
        description = short.get('subText', '')
        tags = "#Shorts #Motivation"  # optional
        playlist = "Inspirational Quotes"  # optional

        save_details_in_excel(video_path, title, description, tags, playlist)

        # Small delay between recordings
        time.sleep(1)

    print("\n🎬 All shorts recorded successfully!")

if __name__ == "__main__":
    main()