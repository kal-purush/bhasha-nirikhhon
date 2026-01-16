import altair as alt
import numpy as np
import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.utils.dataframe import dataframe_to_rows

df = pd.read_excel("Sondage.xlsx")

cols1 = [
    c
    for c in df.columns
    if (c.lower()[:6] != "points" and c.lower()[:8] != "feedback")
]

df_redux = df[cols1[7:]]


ne_se_prononce_pas = [
    "Si oui, lequel ?",
    "De quelle région votre entreprise provient-elle ?",
    "Pouvez-vous détailler ?",
    "Si vous avez déjà investi, quelle est la nature de cet investissement ?",
    "Dans le cas où vous avez choisi de ne pas investir, pouvez-vous nous préciser la raison ?",
    "Pouvez-vous détailler ?2",
    "Le programme d’accompagnement Factory 4.0 et/ou l’investissement associé ont-ils contribué faire monter en compétence certains employés ?",
    "Si oui, pouvez-vous préciser ?",
    "Le programme d’accompagnement Factory 4.0 et/ou l’investissement associé ont-ils eu un impact environnemental dans votre entreprise ?",
]
zero_nan = ["Combien ?", "Combien ?2"]


def fill_df(df):
    for c in df.columns:
        if c in ne_se_prononce_pas:
            df[c] = df[c].replace(np.nan, "Ne se prononce pas", regex=True)
        elif c in zero_nan:
            df[c] = df[c].fillna(0)
    return df


df_redux = fill_df(df_redux)

st.dataframe(df_redux)

questions = st.multiselect("question", df_redux.columns)
questions

vals = df_redux[questions].value_counts()
st.table(vals)
# st.dataframe(df_redux[questions])


# num_questions = [
#     "De quelle région votre entreprise provient-elle ?",
#     "Quel est votre chiffre d’affaire annuel ?",
#     "Combien de personnes employez-vous ?",
#     "Comment qualifieriez-vous votre niveau de connaissances dans l’industrie 4.0, avant d’être accompagné par le programme Factory 4.0 ?",
#     "Que vous a apporté l’accompagnement Factory 4.0 ?",
#     "Dans quelles proportions le projet Factory 4.0 a aidé votre entreprise à faire sa transition vers l’industrie 4.0 ?",
#     "Souhaitez-vous explorer de nouvelles opportunités d’accompagnement/de projet avec l’un des partenaires du programme ?",
#     "Si oui, lequel ?",
#     "À la suite de votre accompagnement par le programme Factory 4.0, pensez-vous investir ? (Définition large incluant infrastructure, équipement, software, formation…).",
#     "Si vous avez déjà investi, quelle est la nature de cet investissement ?",
#     "L’accompagnement Factory 4.0 et/ou l’investissement associé ont-ils contribué à vous fournir un avantage compétitif ?",
#     "Le programme d’accompagnement Factory 4.0 et/ou l’investissement associé ont-ils contribué à pérenniser l’emploi dans votre entreprise ?",
#     "Le programme d’accompagnement Factory 4.0 et/ou l’investissement associé ont-ils contribué à créer de nouveaux emplois ?",
#     "Le programme d’accompagnement Factory 4.0 et/ou l’investissement associé ont-ils contribué faire monter en compétence certains employés ?",
#     "Le programme d’accompagnement Factory 4.0 et/ou l’investissement associé ont-ils eu un impact sur les conditions de travail de vos employés (santé, sécurité, ergonomie) ?",
#     "Le programme d’accompagnement Factory 4.0 et/ou l’investissement associé ont-ils eu un impact environnemental dans votre entreprise ?",
# ]


st.sidebar.markdown("# Enregistrement du fichier")
save_name = st.sidebar.text_input("Nom du fichier", "sondage_nettoye")
export = st.sidebar.button("Enregistrer")

if export:
    wb = openpyxl.Workbook()
    ws = wb.active

    for r in dataframe_to_rows(df_redux, index=True, header=True):
        ws.append(r)

    for cell in ws["A"] + ws[1]:
        cell.style = "Pandas"

    wb.save(f"{save_name}.xlsx")