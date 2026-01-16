from fastapi import APIRouter, HTTPException, Response, Query
from Backend.firebase_config import client
from google.cloud import bigquery
from io import BytesIO
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as XLImage
import pandas as pd
import matplotlib.pyplot as plt

stats_router = APIRouter(prefix="/api/v1/stats", tags=["Statistics"])


@stats_router.get("/top-stores")
def get_top_stores():
    query = """
        SELECT
          JSON_VALUE(data, '$.store_name') AS store_name,
          COUNT(*) AS total_bonuri,
          SUM(CAST(JSON_VALUE(data, '$.total_amount') AS FLOAT64)) AS total_cheltuit
        FROM `receipt-tracking-application.receipts_info.receipts_raw_latest`
        GROUP BY store_name
        ORDER BY total_cheltuit DESC
        LIMIT 5
    """
    try:
        result = client.query(query).result()
        return [
            {
                "store_name": row.store_name,
                "bonuri": row.total_bonuri,
                "total_cheltuit": row.total_cheltuit
            } for row in result
        ]
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Eroare BigQuery: {e}")


@stats_router.get("/distributia-operatiilor-crud")
def get_operation_counts():
    query = """
    SELECT 
      operation,
      COUNT(*) AS count
    FROM `receipt-tracking-application.receipts_info.receipts_raw_changelog`
    GROUP BY operation
    ORDER BY count DESC
    """

    try:
        results = client.query(query).result()
        return [{"operation": row.operation, "count": row.count} for row in results]
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Eroare BigQuery: {e}")


@stats_router.get("/users-spendings")
def get_top_users():
    query = """
    SELECT 
      JSON_VALUE(data, '$.user_uid') AS user_uid,
      SUM(CAST(JSON_VALUE(data, '$.total_amount') AS FLOAT64)) AS total_spent
    FROM `receipt-tracking-application.receipts_info.receipts_raw_latest`
    GROUP BY user_uid
    ORDER BY total_spent DESC
    """

    try:
        results = client.query(query).result()
        rows = [
            {
                "user_uid": row.user_uid,
                "total_spent": row.total_spent
            }
            for row in results
        ]
        return rows
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Eroare BigQuery: {e}")


@stats_router.get("/change-log-info")
def get_change_log():
    query = """
    SELECT *
    FROM `receipt-tracking-application.receipts_info.receipts_raw_changelog`
    ORDER BY timestamp DESC
    """

    try:
        results = client.query(query).result()
        rows = [dict(row.items()) for row in results]
        return rows
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Eroare BigQuery: {e}")


@stats_router.get("/export-excel")
def export_excel_for_user(uid: str = Query(..., description="UID-ul utilizatorului")):
    query = """
    SELECT
      JSON_VALUE(data, '$.store_name') AS store_name,
      CAST(JSON_VALUE(data, '$.total_amount') AS FLOAT64) AS total_amount,
      JSON_VALUE(data, '$.date') AS date
    FROM `receipt-tracking-application.receipts_info.receipts_raw_latest`
    WHERE JSON_VALUE(data, '$.user_uid') = @uid
    ORDER BY date DESC
    """
    job_config = bigquery.QueryJobConfig(
        query_parameters=[bigquery.ScalarQueryParameter("uid", "STRING", uid)]
    )
    job = client.query(query, job_config=job_config)
    results = job.result()
    rows = list(results)

    if not rows:
        raise HTTPException(status_code=404, detail="Nu s-au găsit bonuri pentru acest utilizator.")

    df = pd.DataFrame([dict(r.items()) for r in rows])
    df["date"] = pd.to_datetime(df["date"])
    total_spent = df["total_amount"].sum()
    top_stores = df["store_name"].value_counts().reset_index()
    top_stores.columns = ["store_name", "num_receipts"]

    plt.figure(figsize=(6, 4))
    plt.bar(top_stores["store_name"], top_stores["num_receipts"])
    plt.xticks(rotation=45, ha='right')
    plt.title("Cele mai frecventate magazine")
    plt.tight_layout()
    img_buffer = BytesIO()
    plt.savefig(img_buffer, format='png')
    img_buffer.seek(0)

    wb = Workbook()
    ws_bonuri = wb.active
    ws_bonuri.title = "Bonuri"

    for r in dataframe_to_rows(df, index=False, header=True):
        ws_bonuri.append(r)
    for row in ws_bonuri.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "DD-MM-YYYY"

    ws_stats = wb.create_sheet(title="Statistici")
    ws_stats.append(["Total cheltuit (RON)", round(total_spent, 2)])
    ws_stats.append([])
    ws_stats.append(["Top magazine", "Nr. bonuri"])
    for index, row in top_stores.iterrows():
        ws_stats.append([row["store_name"], row["num_receipts"]])

    img = XLImage(img_buffer)
    ws_stats.add_image(img, "E2")

    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    excel_bytes = excel_buffer.getvalue()
    filename = f"raport_{uid}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename={filename}"}

    return Response(content=excel_bytes, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)